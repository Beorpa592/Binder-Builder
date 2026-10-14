"""Table of Contents workbook (SPEC §8), rendered from an Excel template with `{{ }}` tokens.

Template rows whose first cell starts with `{{#section}}`, `{{#header}}`,
`{{#item}}` or `{{#subitem}}` are row templates: one copy (values and styles) is
written per TOC row of that kind, in tree order, starting where the first row
template sat. Row tokens: mark, code, title, label, page. Any other cell may use
`{{binder_name}}` and `{{page_heading}}` ("Page" in continuous mode, else blank).
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

from openpyxl import load_workbook

from ..tree import Node

_TOKEN = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")
_ROW_MARK = re.compile(r"^\{\{#(section|header|item|subitem)\}\}")


def toc_rows(ctx) -> list[dict]:
    layout = ctx.layout
    rows: list[dict] = []

    def item_rows(nodes: list[Node], depth: int) -> None:
        for n in nodes:
            if not n.is_item:
                continue
            e = layout.entity(n.id)
            rows.append({"kind": "item" if depth == 1 else "subitem", "mark": "X" if e and e.included else "",
                         "code": layout.codes.get(n.id, ""), "title": n.text, "label": "", "page": ""})
            item_rows(n.children, depth + 1)

    def walk(nodes: list[Node]) -> None:
        for n in nodes:
            if n.is_header:
                rows.append({"kind": "header", "mark": "", "code": "", "title": "", "label": n.text, "page": ""})
                walk(n.children)
            elif n.is_section:
                rows.append({"kind": "section", "mark": "", "code": layout.codes[n.id], "title": "",
                             "label": n.text, "page": ""})
                for c in n.children:
                    if c.is_item:
                        item_rows([c], 1)
                    else:
                        walk([c])

    walk(ctx.tree.roots)
    return rows


def _fill(value, vars_: dict):
    if not isinstance(value, str):
        return value, False
    is_code = re.search(r"\{\{\s*code\s*\}\}", value) is not None
    text = _TOKEN.sub(lambda m: str(vars_.get(m.group(1), "")), value)
    return (text if text != "" else None), is_code


def render(ctx, template: Path, out: Path) -> list[str]:
    wb = load_workbook(str(template))
    ws = wb.worksheets[0]
    continuous = ctx.binder.page_numbers == "continuous"
    globals_ = {"binder_name": ctx.binder.name, "page_heading": "Page" if continuous else ""}

    templates: dict[str, tuple[list, float | None]] = {}
    template_rows: list[int] = []
    for row in ws.iter_rows():
        first = row[0]
        m = _ROW_MARK.match(first.value) if isinstance(first.value, str) else None
        if not m:
            continue
        cells = [(c.column, c.value[m.end():] if c is first else c.value, copy.copy(c._style)) for c in row]
        templates[m.group(1)] = (cells, ws.row_dimensions[first.row].height)
        template_rows.append(first.row)
    if "item" not in templates:
        raise ValueError(f"TOC template {template.name} has no {{{{#item}}}} row")

    for row in ws.iter_rows():
        for c in row:
            if c.row not in template_rows and isinstance(c.value, str) and "{{" in c.value:
                c.value = _fill(c.value, globals_)[0]

    start = min(template_rows)
    ws.delete_rows(start, max(template_rows) - start + 1)

    pages = ctx.pages or {}
    r = start
    for row in toc_rows(ctx):
        cells, height = templates.get(row["kind"]) or templates["item"]
        vars_ = dict(globals_, **row)
        vars_["page"] = pages.get(row["code"], "") if continuous else ""
        for col, value, style in cells:
            cell = ws.cell(row=r, column=col)
            cell._style = copy.copy(style)
            v, is_code = _fill(value, vars_)
            cell.value = v
            if is_code and v is not None:
                cell.number_format = "@"
                cell.data_type = "s"
        if height is not None:
            ws.row_dimensions[r].height = height
        r += 1
    wb.save(str(out))
    return []
