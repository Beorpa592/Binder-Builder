"""Master Checklist hybrid rows (SPEC §8.1).

`Reference` (a `=bb_<id>` formula) and `Description` (text) are app-owned, plus a
hidden `bb_id` column used to match rows after the user re-sorts or inserts rows.
Every other column is the user's and is never written. The same mechanism serves
any workbook node marked `hybrid_rows: true`.

Invariant: `user_hash` covers every user-owned cell, row by row, ignoring row
order, so the sync can prove it only touched app-owned cells even after
inserting rows.
"""

from __future__ import annotations

import copy
import hashlib
from dataclasses import dataclass, field

from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from .. import excel

REFERENCE = "Reference"
DESCRIPTION = "Description"
BB_ID = "bb_id"
REMOVED = "[removed]"
NOT_INCLUDED = " [not included]"
REMOVED_RGB = "00FFFF99"
EXCLUDED_RGB = "00D9D9D9"
HEADER_SCAN_ROWS = 20


@dataclass
class RowSpec:
    eid: str
    code: str | None
    title: str
    included: bool = True

    @property
    def reference(self) -> str:
        return f"={excel.name_for(self.eid)}" if self.code else self.title

    @property
    def description(self) -> str:
        return self.title + ("" if self.included else NOT_INCLUDED)


@dataclass
class Changes:
    added: list[str] = field(default_factory=list)
    removed: list[str] = field(default_factory=list)
    updated: list[str] = field(default_factory=list)
    adopted: list[str] = field(default_factory=list)
    appended_at_end: bool = False
    error: str | None = None

    @property
    def any(self) -> bool:
        return bool(self.added or self.removed or self.updated or self.adopted)

    def describe(self) -> str:
        parts = []
        if self.adopted:
            parts.append(f"adopt {len(self.adopted)} legacy row(s)")
        if self.added:
            where = " at the end (sheet has formulas, merges or rules)" if self.appended_at_end else ""
            parts.append(f"add {', '.join(self.added)}{where}")
        if self.removed:
            parts.append(f"mark removed: {', '.join(self.removed)}")
        if self.updated:
            parts.append(f"update {', '.join(self.updated)}")
        return "; ".join(parts)


@dataclass
class SheetLayout:
    ws: object
    header_row: int
    ref_col: int
    desc_col: int
    bb_col: int | None

    @property
    def app_cols(self) -> set[int]:
        return {self.ref_col, self.desc_col} | ({self.bb_col} if self.bb_col else set())


def _norm_heading(v) -> str:
    return str(v).strip().lower() if v is not None else ""


def find_layout(wb, ref_name: str = REFERENCE, desc_name: str = DESCRIPTION) -> SheetLayout | None:
    for ws in excel._user_sheets(wb):
        for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
            heads = {_norm_heading(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1)}
            if ref_name.lower() in heads and desc_name.lower() in heads:
                return SheetLayout(ws, r, heads[ref_name.lower()], heads[desc_name.lower()], heads.get(BB_ID))
    return None


def _last_row(ws) -> int:
    for r in range(ws.max_row, 0, -1):
        if any(c.value is not None for c in ws[r]):
            return r
    return 0


def _safe_to_insert(wb, lay: SheetLayout) -> bool:
    ws = lay.ws
    if ws.merged_cells.ranges or ws.data_validations.dataValidation or len(ws.conditional_formatting) or ws.tables:
        return False
    title = ws.title.lower()
    for sheet in excel._user_sheets(wb):
        for row in sheet.iter_rows():
            for c in row:
                if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                    if sheet is ws and c.column == lay.ref_col and str(c.value).lower().startswith("=bb_"):
                        continue
                    if sheet is ws or title in str(c.value).lower():
                        return False
    for name, dn in wb.defined_names.items():
        if not name.lower().startswith(excel.NAME_PREFIX) and title in (dn.attr_text or "").lower():
            return False
    return True


def _fill_rgb(cell) -> str | None:
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return None
    return f.fgColor.rgb if isinstance(f.fgColor.rgb, str) else None


def _set_fill(cell, rgb: str | None) -> bool:
    current = _fill_rgb(cell)
    if rgb is None:
        if current in (REMOVED_RGB, EXCLUDED_RGB):
            cell.fill = PatternFill(fill_type=None)
            return True
        return False
    if current != rgb:
        cell.fill = PatternFill("solid", fgColor=rgb)
        return True
    return False


def _code_text(v) -> str:
    if isinstance(v, float):
        return f"{v:.2f}"  # legacy CKLs store 0.01 as a float
    return str(v).strip() if v is not None else ""


def sync(wb, specs: list[RowSpec], ref_name: str = REFERENCE, desc_name: str = DESCRIPTION) -> Changes:
    ch = Changes()
    lay = find_layout(wb, ref_name, desc_name)
    if lay is None:
        ch.error = f"no header row with {ref_name!r} and {desc_name!r}; checklist rows not maintained"
        return ch
    ws = lay.ws
    by_id = {s.eid: s for s in specs}

    if lay.bb_col is None:
        lay.bb_col = ws.max_column + 1
        ws.cell(lay.header_row, lay.bb_col, BB_ID)
        ws.column_dimensions[get_column_letter(lay.bb_col)].hidden = True
        by_code = {s.code: s.eid for s in specs if s.code}
        for r in range(lay.header_row + 1, _last_row(ws) + 1):
            eid = by_code.pop(_code_text(ws.cell(r, lay.ref_col).value), None)
            if eid:
                ws.cell(r, lay.bb_col, eid)
                ch.adopted.append(eid)

    rows: dict[str, int] = {}
    for r in range(lay.header_row + 1, _last_row(ws) + 1):
        v = ws.cell(r, lay.bb_col).value
        if v:
            rows.setdefault(str(v), r)

    for eid, r in rows.items():
        ref, desc = ws.cell(r, lay.ref_col), ws.cell(r, lay.desc_col)
        s = by_id.get(eid)
        if s is None:
            changed = ref.value != REMOVED
            ref.value = REMOVED
            changed |= _set_fill(ref, REMOVED_RGB) | _set_fill(desc, REMOVED_RGB)
            if changed:
                ch.removed.append(eid)
            continue
        changed = False
        if ref.value != s.reference:
            ref.value = s.reference
            changed = True
        if desc.value != s.description:
            desc.value = s.description
            changed = True
        rgb = None if s.included else EXCLUDED_RGB
        changed |= _set_fill(ref, rgb) | _set_fill(desc, rgb)
        if changed and eid not in ch.adopted:
            ch.updated.append(eid)

    missing = [s for s in specs if s.eid not in rows]
    if missing:
        safe = _safe_to_insert(wb, lay)
        ch.appended_at_end = not safe
        order = [s.eid for s in specs]
        last = max(_last_row(ws), lay.header_row)
        for s in missing:
            if safe:
                prev = next((rows[e] for e in reversed(order[:order.index(s.eid)]) if e in rows), lay.header_row)
                at = prev + 1
                ws.insert_rows(at)
                for k, v in rows.items():
                    if v >= at:
                        rows[k] = v + 1
                style_row = at - 1 if at - 1 > lay.header_row else at + 1
            else:
                at = last + 1
                style_row = last if last > lay.header_row else None
            last += 1
            if style_row and style_row <= ws.max_row:
                for c in range(1, ws.max_column + 1):
                    src = ws.cell(style_row, c)
                    if src.has_style:
                        ws.cell(at, c)._style = copy.copy(src._style)
            ws.cell(at, lay.bb_col, s.eid)
            ws.cell(at, lay.ref_col, s.reference)
            ws.cell(at, lay.desc_col, s.description)
            rgb = None if s.included else EXCLUDED_RGB
            _set_fill(ws.cell(at, lay.ref_col), rgb)
            _set_fill(ws.cell(at, lay.desc_col), rgb)
            rows[s.eid] = at
            ch.added.append(s.eid)
    return ch


def user_hash(wb, ref_name: str = REFERENCE, desc_name: str = DESCRIPTION) -> str:
    """Order-independent hash of every user-owned cell (see module docstring)."""
    lay = find_layout(wb, ref_name, desc_name)
    if lay is None:
        return excel.body_hash_wb(wb)
    ws = lay.ws
    app = lay.app_cols
    rows = []
    for row in ws.iter_rows():
        vals = tuple((c.column, excel._norm(c.value)) for c in row if c.column not in app and c.value is not None)
        if vals:
            rows.append(repr(vals))
    h = hashlib.sha256()
    h.update(ws.title.encode())
    for r in sorted(rows):
        h.update(r.encode())
        h.update(b"\x1e")
    h.update(excel.body_hash_wb(wb, skip=(ws.title,)).encode())
    return h.hexdigest()
