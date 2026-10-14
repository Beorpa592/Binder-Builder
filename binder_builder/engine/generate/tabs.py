"""Tabs sheet (SPEC §8): one label per section and sub-section, filling the template table's cells.

The template's first table is the layout. Its cells are filled in reading order,
and copies of the pristine table are appended (each after a page break) when
there are more sections than cells. Labels and the fit warning come from
`codes.tab_label`.
"""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.table import Table


def labels(ctx) -> list[str]:
    return [s.tab_label for s in ctx.layout.sections if s.tab_label]


def _unique_cells(table: Table) -> list:
    seen, out = set(), []
    for row in table.rows:
        for cell in row.cells:
            if id(cell._tc) not in seen:
                seen.add(id(cell._tc))
                out.append(cell)
    return out


def _set_cell_text(cell, text: str) -> None:
    """Set a cell's text, keeping the first run's formatting."""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def render(ctx, template: Path, out: Path) -> list[str]:
    doc = Document(str(template))
    if not doc.tables:
        raise ValueError(f"tabs template {template.name} has no table")
    first = doc.tables[0]
    pristine = copy.deepcopy(first._tbl)
    per_table = len(_unique_cells(first))
    items = labels(ctx)
    chunks = [items[i:i + per_table] for i in range(0, max(len(items), 1), per_table)]
    anchor = first._tbl
    for k, chunk in enumerate(chunks):
        if k == 0:
            table = first
        else:
            brk = doc.add_paragraph()
            brk.add_run().add_break(WD_BREAK.PAGE)
            anchor.addnext(brk._p)
            new = copy.deepcopy(pristine)
            brk._p.addnext(new)
            anchor = new
            table = Table(new, first._parent)
        for i, cell in enumerate(_unique_cells(table)):
            _set_cell_text(cell, chunk[i] if i < len(chunk) else "")
    doc.save(str(out))
    return []
