"""Excel (.xlsx) support: header writer, `_binder` sheet, defined names, body hash (SPEC §5.2, §6.2, §7.4)."""

from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from .codes import HeaderSpec

BINDER_SHEET = "_binder"
NAME_PREFIX = "bb_"
BROKEN = "[?]"
HEADER_ROW = ("id", "code", "title", "filename")

# Parts openpyxl cannot round-trip. A workbook containing any of these is never
# written by the app (it would silently lose charts, images, pivots or macros).
_UNSAFE = re.compile(
    r"^xl/(drawings/drawing\d+\.xml|charts/|chartsheets/|pivotTables/|pivotCache/|media/|embeddings/"
    r"|vbaProject\.bin|slicers/|slicerCaches/|timelines/|timelineCaches/|ctrlProps/|activeX/)"
)
_NAME_REF = re.compile(r"(?<![A-Za-z0-9_.])(bb_[A-Za-z0-9_]+)", re.IGNORECASE)


def name_for(node_id: str) -> str:
    return NAME_PREFIX + node_id.replace("-", "_")


def unsafe_parts(path: Path | str) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return sorted({m.group(0) for n in z.namelist() if (m := _UNSAFE.match(n))})


def load(path: Path | str):
    return load_workbook(str(path), rich_text=True)


def _user_sheets(wb):
    return [ws for ws in wb.worksheets if ws.title != BINDER_SHEET]


# ---- body hash (SPEC §7.4) --------------------------------------------------

def _norm(value) -> str:
    if hasattr(value, "text") and hasattr(value, "ref"):  # ArrayFormula
        return f"AF:{value.ref}:{value.text}"
    if type(value).__name__ == "DataTableFormula":
        return f"DT:{sorted(vars(value).items())}"
    if type(value).__name__ == "CellRichText":
        return f"RT:{value!s}"
    return f"{type(value).__name__}:{value!r}"


def body_hash_wb(wb) -> str:
    h = hashlib.sha256()
    # Stricter than SPEC §7.4: hidden user sheets are hashed too (see DEVIATIONS.md).
    for ws in _user_sheets(wb):
        h.update(f"\x1d{ws.title}\x1d".encode())
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    h.update(f"{c.coordinate}={_norm(c.value)}\x1f".encode())
        for name, dn in sorted(ws.defined_names.items()):
            if not name.lower().startswith(NAME_PREFIX):
                h.update(f"L:{name}={dn.attr_text}\x1f".encode())
    for name, dn in sorted(wb.defined_names.items()):
        if not name.lower().startswith(NAME_PREFIX):
            h.update(f"N:{name}={dn.attr_text}\x1f".encode())
    return h.hexdigest()


def body_hash(path: Path | str) -> str:
    return body_hash_wb(load(path))


# ---- headers (SPEC §5.2) ----------------------------------------------------

def _esc(s: str) -> str:
    return s.replace("&", "&&")


def desired_header(spec: HeaderSpec) -> tuple[tuple, tuple]:
    center = "\n".join(_esc(x) for x in spec.lines) or None
    right_parts = []
    if spec.code is not None:
        right_parts.append(_esc(f"{spec.ref_label} {spec.code}"))
    if spec.page_numbers:
        right_parts.append("Page &P of &N")
    right = "\n".join(right_parts) or None
    left_f, center_f, right_f = (_esc(x) or None for x in spec.footer)
    return (None, center, right), (left_f, center_f, right_f)


def _parts(hf) -> tuple:
    return tuple((p.text or None) for p in (hf.left, hf.center, hf.right))


def _header_targets(ws):
    pairs = [(ws.oddHeader, ws.oddFooter)]
    if ws.HeaderFooter.differentOddEven:
        pairs.append((ws.evenHeader, ws.evenFooter))
    if ws.HeaderFooter.differentFirst:
        pairs.append((ws.firstHeader, ws.firstFooter))
    return pairs


def header_matches(wb, spec: HeaderSpec) -> bool:
    want_h, want_f = desired_header(spec)
    for ws in _user_sheets(wb):
        for hdr, ftr in _header_targets(ws):
            if _parts(hdr) != want_h or _parts(ftr) != want_f:
                return False
    return True


def write_header(wb, spec: HeaderSpec) -> None:
    want_h, want_f = desired_header(spec)
    for ws in _user_sheets(wb):
        for hdr, ftr in _header_targets(ws):
            for part, text in zip((hdr.left, hdr.center, hdr.right), want_h):
                part.text = text
            for part, text in zip((ftr.left, ftr.center, ftr.right), want_f):
                part.text = text


def read_code(path: Path | str) -> str | None:
    """The code from the first sheet's right header block (for tests)."""
    wb = load(path)
    right = _user_sheets(wb)[0].oddHeader.right.text or ""
    m = re.search(r"#:?\s*(\S+)", right.split("\n")[0])
    return m.group(1) if m else None


# ---- _binder sheet and names (SPEC §6.2) --------------------------------------

@dataclass
class BinderRows:
    rows: list[tuple[str, str, str, str]]  # id, code, title, filename

    @property
    def names(self) -> dict[str, str]:
        return {name_for(r[0]): f"'{BINDER_SHEET}'!$B${i}" for i, r in enumerate(self.rows, start=2)}


def referenced_names(wb) -> dict[str, list[str]]:
    """bb_ names used in formulas -> list of 'Sheet!A1' locations."""
    found: dict[str, list[str]] = {}
    for ws in _user_sheets(wb):
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                text = v.text if hasattr(v, "text") and hasattr(v, "ref") else v
                if isinstance(text, str) and text.startswith("=") or c.data_type == "f":
                    for m in _NAME_REF.finditer(str(text)):
                        found.setdefault(m.group(1).lower(), []).append(f"{ws.title}!{c.coordinate}")
    return found


def read_binder_rows(wb) -> list[tuple[str, str, str, str]] | None:
    if BINDER_SHEET not in wb.sheetnames:
        return None
    ws = wb[BINDER_SHEET]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None:
            continue
        rows.append(tuple("" if v is None else str(v) for v in (list(r) + [None] * 4)[:4]))
    return rows


def full_rows(wb, rows: list[tuple[str, str, str, str]]) -> tuple[BinderRows, dict[str, list[str]]]:
    """Add `[?]` rows for names used in formulas whose target no longer exists.

    Returns the rows to write and {broken name: [cell locations]}.
    """
    known = {name_for(r[0]) for r in rows}
    old_ids = {name_for(r[0]): r[0] for r in (read_binder_rows(wb) or [])}
    broken: dict[str, list[str]] = {}
    extra = []
    for name, cells in sorted(referenced_names(wb).items()):
        if name in known:
            continue
        broken[name] = cells
        node_id = old_ids.get(name) or name[len(NAME_PREFIX):].replace("_", "-")
        extra.append((node_id, BROKEN, "", ""))
    return BinderRows(list(rows) + extra), broken


def binder_sheet_matches(wb, want: BinderRows) -> bool:
    if BINDER_SHEET not in wb.sheetnames:
        return False
    if wb[BINDER_SHEET].sheet_state != "hidden":
        return False
    if read_binder_rows(wb) != want.rows:
        return False
    have = {n: dn.attr_text for n, dn in wb.defined_names.items() if n.lower().startswith(NAME_PREFIX)}
    return have == want.names


def write_binder_sheet(wb, want: BinderRows) -> None:
    if BINDER_SHEET in wb.sheetnames:
        ws = wb[BINDER_SHEET]
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(BINDER_SHEET)
    for r_idx, row in enumerate([HEADER_ROW] + list(want.rows), start=1):
        for c_idx, v in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=str(v))
            cell.number_format = "@"
            cell.data_type = "s"
    ws.sheet_state = "hidden"
    for n in [n for n in wb.defined_names if n.lower().startswith(NAME_PREFIX)]:
        del wb.defined_names[n]
    for n, ref in want.names.items():
        wb.defined_names[n] = DefinedName(n, attr_text=ref)
    # A hidden sheet must never be the active one.
    if wb.active is not None and wb.active.title == BINDER_SHEET:
        wb.active = 0
    for ws_ in wb.worksheets:
        if ws_.title == BINDER_SHEET:
            ws_.sheet_view.tabSelected = False


def resolve_name(path: Path | str, name: str) -> str | None:
    """Value of the cell a bb_ name points at (for tests)."""
    wb = load(path)
    dn = wb.defined_names.get(name)
    if dn is None:
        return None
    for sheet, coord in dn.destinations:
        return wb[sheet][coord.replace("$", "")].value
    return None


# ---- file-level operations used by sync -------------------------------------

@dataclass
class BookState:
    header_ok: bool
    refs_ok: bool
    broken: dict[str, list[str]] = field(default_factory=dict)


def inspect(path: Path | str, spec: HeaderSpec | None, rows: list[tuple[str, str, str, str]]) -> BookState:
    wb = load(path)
    ok = True if spec is None else header_matches(wb, spec)
    want, broken = full_rows(wb, rows)
    return BookState(ok, binder_sheet_matches(wb, want), broken)


def update(src: Path | str, dest: Path | str, spec: HeaderSpec | None,
           rows: list[tuple[str, str, str, str]]) -> None:
    wb = load(src)
    if spec is not None and not header_matches(wb, spec):
        write_header(wb, spec)
    want, _ = full_rows(wb, rows)
    if not binder_sheet_matches(wb, want):
        write_binder_sheet(wb, want)
    wb.save(str(dest))


_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def sheet_title(title: str) -> str:
    s = _BAD_SHEET_CHARS.sub("", title).strip("'").strip() or "Sheet1"
    return s[:31]


def create_blank(path: Path | str, title: str) -> None:
    wb = Workbook()
    wb.active.title = sheet_title(title)
    wb.save(str(path))
