"""Rebuild the default template set shipped with the app.

    python -m binder_builder.templates.build_defaults

Staff restyle a binder by copying these into the binder's template folder
(`.binder/templates/` unless `binder.template_folder` says otherwise) and editing
them in Word or Excel. Placeholders: docxtpl (Jinja) in Word, `{{ }}` tokens in
Excel cells. See templates/README.md.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parent


def build_cover(out: Path) -> None:
    doc = Document()
    for _ in range(6):
        doc.add_paragraph()
    doc.add_paragraph("{%p for line in lines %}")
    p = doc.add_paragraph("{{ line }}", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("{%p endfor %}")
    doc.save(str(out))


def build_section_cover(out: Path) -> None:
    doc = Document()
    doc.add_paragraph("{{ heading }}", style="Title")
    doc.add_paragraph("{%p if description %}")
    doc.add_paragraph("{{ description }}", style="Subtitle")
    doc.add_paragraph("{%p endif %}")
    doc.add_paragraph()
    doc.add_paragraph("{%p for r in rows %}")
    p = doc.add_paragraph("{{ r.mark }}\t{{ r.code }}\t{{ r.title }}")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.9))
    p.paragraph_format.tab_stops.add_tab_stop(Inches(2.0))
    p.paragraph_format.space_after = Pt(6)
    doc.add_paragraph("{%p endfor %}")
    doc.save(str(out))


def build_tabs(out: Path, cells: int = 8) -> None:
    doc = Document()
    t = doc.add_table(rows=cells, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for row in t.rows:
        row.height = Inches(1.0)
        p = row.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("")
        run.font.size = Pt(16)
        run.font.bold = True
    doc.save(str(out))


def build_toc(out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "TOC"
    bold = Font(bold=True)
    ws["A1"] = "{{binder_name}}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Table of Contents"
    ws["A2"].font = Font(bold=True, size=12)
    for col, text in zip("ABCDE", ("Incl.", "Ref", "Title", "", "{{page_heading}}")):
        ws[f"{col}3"] = text or None
        ws[f"{col}3"].font = bold
    # Row templates: a row whose first cell starts with {{#kind}} is repeated once per TOC row of that kind.
    ws.append(["{{#section}}", "{{code}}", "{{label}}"])
    ws.append(["{{#header}}", None, "{{label}}"])
    ws.append(["{{#item}}{{mark}}", "{{code}}", "{{title}}", None, "{{page}}"])
    ws.append(["{{#subitem}}{{mark}}", None, "{{code}}", "{{title}}", "{{page}}"])
    ws["B4"].font = bold
    ws["C4"].font = bold
    ws["C5"].font = Font(bold=True, italic=True)
    for cell in ("B4", "B6", "C7"):
        ws[cell].number_format = "@"
    ws["A6"].alignment = Alignment(horizontal="center")
    ws["A7"].alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDE", (6, 10, 48, 40, 8)):
        ws.column_dimensions[col].width = width
    wb.save(str(out))


def build_ckl(out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Checklist"
    ws.append(["Reference", "Description", "Hours", "Reviewer", "Date Reviewed", "Final Review"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, width in zip("ABCDEF", (12, 48, 8, 20, 14, 14)):
        ws.column_dimensions[col].width = width
    wb.save(str(out))


def build_all(folder: Path = HERE) -> None:
    build_cover(folder / "00_Cover.docx")
    build_section_cover(folder / "Section_Cover.docx")
    build_tabs(folder / "Tabs.docx")
    build_toc(folder / "01_TOC.xlsx")
    build_ckl(folder / "02_CKL.xlsx")


if __name__ == "__main__":
    build_all()
    print(f"templates written to {HERE}")
