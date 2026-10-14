"""Phase 0 acceptance tests (SPEC §14) plus the safety rules of §1.1."""

import json
import subprocess
import sys

import pytest

from binder_builder.engine import excel, sync, word
from binder_builder.engine.manifest import Manifest
from binder_builder.engine.storage import LocalFolderStorage, LockedError

from .helpers import body_hashes_by_entity, do_sync, edit_tree, office_files

ADMIN = "A_Administration"


def _all_files(folder):
    return {p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file() and ".binder" not in p.parts}


# ---- acceptance: load, sync, empty second plan ---------------------------------

def test_first_sync_then_empty_plan(binder):
    before_files = _all_files(binder)
    before_hashes = {p.relative_to(binder).as_posix(): sync.body_hash(p) for p in office_files(binder)}

    p1 = sync.plan(binder)
    kinds = {a.kind for a in p1.changes}
    assert sync.STAMP in kinds  # legacy headers get taken over
    assert sync.REFS in kinds  # _binder sheets get written
    assert not p1.orphans
    assert not p1.warnings, [a.detail for a in p1.warnings]
    do_sync(binder)

    p2 = sync.plan(binder)
    assert p2.is_empty, p2.render()
    assert not p2.actions, p2.render()

    # nothing moved, created or lost; no body changed
    assert _all_files(binder) == before_files
    for rel, h in before_hashes.items():
        assert sync.body_hash(binder / rel) == h, rel


def test_first_sync_headers_uniform_and_locked_untouched(binder):
    locked = binder / "B_Policies_and_Procedures/03_Licensed_Release_Form.docx"
    locked_before = word.header_text(locked)
    do_sync(binder)
    assert word.read_code(binder / ADMIN / "03_Top_10_List.docx") == "A.03"  # was labelled "Ref #:"
    assert word.read_code(binder / ADMIN / "03_PB_Top_10_List_Backup.docx") == "A.03.PB"
    assert word.read_code(binder / "A_Administration/00_Cover.docx") == "A.00"
    assert excel.read_code(binder / "0_Control_Section/01_Open_Points.xlsx") == "0.01"
    assert excel.read_code(binder / "01_TOC_m.xlsx") == "TOC"
    assert word.header_text(locked) == locked_before
    # binder cover and tabs are not stamped
    assert word.read_code(binder / "00_Cover.docx") is None


# ---- acceptance: insert between A.02 and A.03 ---------------------------------

EXPECTED_AFTER_INSERT = {
    "press-kit": ("A.03", "03_Press_Kit.docx"),
    "top-10-list": ("A.04", "04_Top_10_List.docx"),
    "top-10-list-backup": ("A.04.PB", "04_PB_Top_10_List_Backup.docx"),
    "points-of-contact": ("A.05", "05_Points_of_Contact.docx"),
    "budget-overview": ("A.06", "06_Budget_Overview.xlsx"),
    "meeting-minutes": ("A.07", "07_Meeting_Minutes.docx"),
    "outreach-efforts": ("A.08", "08_Social_Media_Outreach_Efforts.docx"),
}


def test_insert_item_renumbers_everything(binder):
    do_sync(binder)
    hashes_before = body_hashes_by_entity(binder)

    m, tree = edit_tree(binder)
    tree.insert_after("org-chart", {"kind": "item", "title": "Press Kit", "ext": "docx"})
    m.save()

    p = sync.plan(binder)
    moves = {a.entity_id: a.target for a in p.changes if a.kind == sync.MOVE}
    assert set(moves) == set(EXPECTED_AFTER_INSERT) - {"press-kit"}
    assert [a.path for a in p.changes if a.kind == sync.CREATE] == [f"{ADMIN}/03_Press_Kit.docx"]
    r = sync.apply(p)
    assert r.ok, r.failures

    # every file from A.03 onward renamed, old names gone
    for eid, (code, name) in EXPECTED_AFTER_INSERT.items():
        path = binder / ADMIN / name
        assert path.is_file(), name
        reader = word.read_code if path.suffix == ".docx" else excel.read_code
        assert reader(path) == code, (eid, reader(path))
    for old in ("03_Top_10_List.docx", "03_PB_Top_10_List_Backup.docx", "04_Points_of_Contact.docx",
                "05_Budget_Overview.xlsx", "06_Meeting_Minutes.docx", "07_Social_Media_Outreach_Efforts.docx"):
        assert not (binder / ADMIN / old).exists(), old
    # untouched positions keep their names and codes
    assert word.read_code(binder / ADMIN / "02_Organizational_Chart.docx") == "A.02"

    # every bb:ref control shows the new code
    from docx import Document
    mrp = {h.target_id: h.current for h in word.scan_refs(Document(str(binder / ADMIN / "01_North_Star_MRP_Summary.docx")), {})}
    assert mrp == {"top-10-list": "A.04", "outreach-efforts": "A.08"}
    out = {h.target_id: h.current for h in word.scan_refs(Document(str(binder / ADMIN / "08_Social_Media_Outreach_Efforts.docx")), {})}
    assert out == {"top-10-list-backup": "A.04.PB"}

    # every bb_ name in every workbook resolves to the new code
    for xlsx in [x for x in office_files(binder) if x.suffix == ".xlsx"]:
        assert excel.resolve_name(xlsx, "bb_top_10_list") == "A.04", xlsx
        assert excel.resolve_name(xlsx, "bb_outreach_efforts") == "A.08", xlsx
        assert excel.resolve_name(xlsx, "bb_press_kit") == "A.03", xlsx
        assert excel.resolve_name(xlsx, "bb_mrp_summary") == "A.01", xlsx

    # every body hash unchanged
    hashes_after = body_hashes_by_entity(binder)
    for eid, h in hashes_before.items():
        assert hashes_after[eid] == h, eid

    assert sync.plan(binder).is_empty


# ---- acceptance: delete a node ------------------------------------------------

def test_removed_node_file_stays_and_is_orphaned(binder):
    do_sync(binder)
    m, tree = edit_tree(binder)
    tree.remove("points-of-contact")
    m.save()
    p = sync.plan(binder)
    assert [a.path for a in p.orphans] == [f"{ADMIN}/04_Points_of_Contact.docx"]
    r = sync.apply(p)
    assert r.ok
    assert (binder / ADMIN / "04_Points_of_Contact.docx").is_file()
    assert (binder / ADMIN / "04_Budget_Overview.xlsx").is_file()  # A.05 moved up to A.04
    p2 = sync.plan(binder)
    assert p2.is_empty
    assert [a.path for a in p2.orphans] == [f"{ADMIN}/04_Points_of_Contact.docx"]


def test_removed_referenced_node_breaks_refs_visibly(binder):
    do_sync(binder)
    m, tree = edit_tree(binder)
    tree.remove("top-10-list")
    m.save()
    p = sync.plan(binder)
    broken = [a for a in p.warnings if "broken reference" in a.detail]
    assert any("Top 10 List" in a.detail and a.path.endswith("01_North_Star_MRP_Summary.docx") for a in broken)
    assert any("bb_top_10_list" in a.detail and "Open_Points" in a.path for a in broken)
    do_sync(binder)
    from docx import Document
    mrp = {h.target_id: h.current for h in word.scan_refs(Document(str(binder / ADMIN / "01_North_Star_MRP_Summary.docx")), {})}
    assert mrp["top-10-list"] == "[?]"
    assert excel.resolve_name(binder / "0_Control_Section/01_Open_Points.xlsx", "bb_top_10_list") == "[?]"
    assert (binder / ADMIN / "03_Top_10_List.docx").is_file()


# ---- safety rules -----------------------------------------------------------------

def test_never_deletes_files_across_reorders(binder):
    do_sync(binder)
    m, tree = edit_tree(binder)
    tree.move("policies", None, 1)  # Policies before Administration: A <-> B swap
    m.save()
    before = len(_all_files(binder))
    do_sync(binder)
    assert len(_all_files(binder)) == before
    assert (binder / "A_Policies_and_Procedures/01_Media_Policy.docx").is_file()
    assert word.read_code(binder / "A_Policies_and_Procedures/01_Media_Policy.docx") == "A.01"
    assert word.read_code(binder / "B_Administration/03_PB_Top_10_List_Backup.docx") == "B.03.PB"
    assert not (binder / "A_Administration").exists()  # emptied folder tidied, no files lost
    assert sync.plan(binder).is_empty


def test_backup_snapshots_every_touched_file(binder):
    r = do_sync(binder)
    backup = binder / ".binder" / "backups" / r.backup
    index = json.loads((backup / "index.json").read_text())
    assert index["files"]
    for entry in index["files"]:
        assert (backup / entry["path"]).is_file()
        # the snapshot is the original bytes
        assert sync.file_hash(backup / entry["path"]) == sync.file_hash(
            __import__("pathlib").Path(__file__).parent / "fixtures" / "synthetic_binder" / entry["path"])


def test_invariant_failure_keeps_original(binder, monkeypatch):
    target = binder / ADMIN / "02_Organizational_Chart.docx"
    original = sync.file_hash(target)
    real_update = word.update

    def evil_update(src, dest, spec, codes):
        real_update(src, dest, spec, codes)
        from docx import Document
        d = Document(str(dest))
        d.paragraphs[1].text = "tampered"
        d.save(str(dest))

    monkeypatch.setattr(word, "update", evil_update)
    r = sync.apply(sync.plan(binder))
    assert not r.ok
    assert any("invariant" in f for f in r.failures)
    assert sync.file_hash(target) == original


def test_restore_puts_files_back(binder):
    do_sync(binder)
    m, tree = edit_tree(binder)
    tree.insert_after("org-chart", {"kind": "item", "title": "Press Kit", "ext": "docx"})
    m.save()
    snapshot = {rel: sync.file_hash(binder / rel) for rel in _all_files(binder)}
    r = do_sync(binder)
    restored = sync.restore(binder, r.backup)
    assert restored
    for rel, h in snapshot.items():
        assert sync.file_hash(binder / rel) == h, rel
    assert (binder / ADMIN / "03_Press_Kit.docx").is_file()  # created files are never deleted


def test_lock_held_by_someone_else_blocks_apply(binder):
    lock = binder / ".binder" / "binder.lock"
    lock.write_text(json.dumps({"user": "someone", "machine": "DESK-07", "since": "2026-10-14T15:30:00"}))
    with pytest.raises(LockedError) as e:
        sync.apply(sync.plan(binder))
    assert "Locked by someone on DESK-07" in str(e.value)
    assert lock.exists()


def test_lock_released_after_sync(binder):
    do_sync(binder)
    assert LocalFolderStorage(binder).lock_info() is None
    log = LocalFolderStorage(binder).history()
    assert log and log[-1]["actions"] > 0


def test_pdf_item_without_file_warns(binder):
    (binder / "B_Policies_and_Procedures/02_Signed_Policy_Acknowledgement.pdf").rename(binder / "elsewhere.pdf")
    p = sync.plan(binder)
    assert any("needs a file" in a.detail for a in p.warnings)
    assert any(a.path == "elsewhere.pdf" for a in p.orphans)


def test_included_false_keeps_code_and_file(binder):
    do_sync(binder)
    assert word.read_code(binder / "B_Policies_and_Procedures/04_Retired_Policy.docx") == "B.04"


def test_manifest_comments_survive_edit(binder):
    m, tree = edit_tree(binder)
    tree.insert_after("org-chart", {"kind": "item", "title": "Press Kit", "ext": "docx"})
    m.save()
    text = (binder / ".binder" / "binder.yaml").read_text()
    assert "# clarifying header: no letter, no folder, no tab" in text
    assert "press-kit" in text
    assert Manifest.load(binder).model.tree[1].children[2].id == "press-kit"


def test_cli_plan_and_apply(binder):
    run = lambda *a: subprocess.run([sys.executable, "-m", "binder_builder", *a], capture_output=True, text=True)
    out = run("sync", str(binder), "--plan")
    assert out.returncode == 0, out.stderr
    assert "stamp_header" in out.stdout
    out = run("sync", str(binder), "--apply", "--yes")
    assert out.returncode == 0, out.stderr
    out = run("sync", str(binder), "--plan")
    assert "Nothing to do" in out.stdout
