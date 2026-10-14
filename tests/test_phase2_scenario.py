"""Phase 2 acceptance scenario (SPEC §14), driven through the same session the GUI uses.

"A non-technical user creates a binder with five sections and twelve documents,
syncs, opens the files in Word, reorders two sections, syncs again, and the
files, headers, TOC, covers, and tabs are correct, with zero body changes."

Opening in Word is simulated by writing body content the way an author would.
`test_office_smoke.py` opens the same kind of files in real Word/Excel.
"""

from docx import Document
from openpyxl import load_workbook

from binder_builder.engine import excel, sync, word
from binder_builder.engine.session import BinderSession

from .dump import dump

STRUCTURE = {
    "Administration": [("Org Chart", "docx"), ("Budget", "xlsx"), ("Contacts", "docx")],
    "Policies": [("Media Policy", "docx"), ("Privacy Policy", "docx")],
    "Campaigns": [("Calendar", "xlsx"), ("Briefs", "docx"), ("Results", "xlsx")],
    "Website": [("Site Map", "docx"), ("Analytics", "xlsx")],
    "Print": [("Inventory", "xlsx"), ("Flyers", "docx")],
}


def _author_writes(path):
    """What an author does in Word/Excel: body text only."""
    if path.suffix == ".docx":
        d = Document(str(path))
        d.add_paragraph(f"Author text in {path.stem}. Placeholder content.")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "cell"
        d.save(str(path))
    else:
        wb = load_workbook(path)
        ws = wb.worksheets[0]
        ws["A1"] = f"Author data in {path.stem}"
        ws["B2"] = 42
        wb.save(path)


def test_phase2_scenario(tmp_path):
    folder = tmp_path / "Scenario Binder"
    s = BinderSession.create(folder, "Scenario Binder")
    s.update_binder(header_lines=["{{org_short}}", "Scenario Binder", "{{doc_title}}"])
    prev = None
    for label, docs in STRUCTURE.items():
        sec = s.add_section(label, prev, "after")
        last = None
        for title, ext in docs:
            node = s.add_document(title, ext, last or sec.id, "after" if last else "child")
            last = node.id
        prev = sec.id
    layout = s.layout()
    assert len(layout.sections) == 5 and len(layout.items()) == 12

    r = s.apply(s.plan())
    assert r.ok, r.failures
    assert s.plan().is_empty
    for e in layout.items():
        assert (folder / e.rel_path).is_file()

    # authors open the files and write their content
    for e in layout.items():
        _author_writes(folder / e.rel_path)
    s.apply(s.plan())  # nothing app-owned changed, but this proves an author edit doesn't trigger rewrites
    hashes = {e.entity_id: sync.body_hash(folder / e.rel_path) for e in s.layout().items()}

    # reorder two sections: Print moves to the top, Policies moves to the end
    s.move("print", None, 0)
    s.move("policies", None, len(s.tree.roots))
    r = s.apply(s.plan())
    assert r.ok, r.failures
    assert s.plan().is_empty

    layout = s.layout()
    assert [sec.node.id for sec in layout.sections] == ["print", "administration", "campaigns", "website", "policies"]
    assert layout.codes["inventory"] == "A.01" and layout.codes["media-policy"] == "E.01"
    for e in layout.items():
        p = folder / e.rel_path
        assert p.is_file(), e.rel_path
        reader = word.read_code if p.suffix == ".docx" else excel.read_code
        assert reader(p) == e.code, e.rel_path
        assert sync.body_hash(p) == hashes[e.entity_id], f"body changed: {e.rel_path}"
    assert not (folder / "B_Policies").exists() and (folder / "E_Policies" / "01_Media_Policy.docx").is_file()

    toc = dump(folder / "01_TOC.xlsx")
    assert "B4='A' @  C4='Print'" in toc
    assert toc.index("'Print'") < toc.index("'Administration'") < toc.index("'Policies'")
    assert "A) Print" in dump(folder / "A_Print" / "00_Cover.docx")
    assert "E) Policies" in dump(folder / "E_Policies" / "00_Cover.docx")
    tabs = dump(folder / "Tabs.docx")
    assert tabs.index("A) Print") < tabs.index("B) Administration") < tabs.index("E) Policies")
    assert not sync.plan(folder).orphans
    s.close()
