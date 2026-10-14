"""Phase 3 acceptance tests (SPEC §14): import, claim, fill.

"A legacy-style fixture (floats in the TOC, inconsistent header labels, one stale
code) imports with every disagreement shown and nothing changed until accepted;
after first sync all headers are uniform and claimed references resolve."
"""

import json
import shutil
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from binder_builder.engine import excel, importer, refs, sync, word
from binder_builder.engine.codes import compute_layout
from binder_builder.engine.manifest import Manifest
from binder_builder.engine.session import BinderSession
from binder_builder.engine.tree import Tree

LEGACY = Path(__file__).parent / "fixtures" / "legacy_binder"
ADMIN = "A_Administration"


@pytest.fixture
def legacy(tmp_path) -> Path:
    dest = tmp_path / "legacy" / "North Star Media Management Binder"
    shutil.copytree(LEGACY, dest)
    return dest


def _snapshot(folder: Path) -> dict[str, str]:
    return {p.relative_to(folder).as_posix(): sync.file_hash(p) for p in folder.rglob("*") if p.is_file()}


# ---- TOC parsing -------------------------------------------------------------------------

def test_toc_floats_become_string_codes_and_structure_is_rebuilt(legacy):
    rows = importer.read_toc(legacy / "01_TOC_m.xlsx")
    codes = [r.code for r in rows if r.code]
    assert "0.01" in codes and "0" in codes and all(isinstance(c, str) for c in codes)
    plan = importer.analyze(legacy)
    tree = Tree(plan.manifest)
    assert tree.find("control-section").get("numeric") is True
    assert tree.find("retired-policy").included is False
    assert tree.find("top-10-list-backup").get("suffix") == "PB"
    assert tree.find("weekly-digest-monthly-summary").parent.id == "index-of-genres-and-channels"
    header = tree.find("content-mediums")
    assert header.kind == "header" and [c.id for c in header.children] == ["website-media", "print-media"]
    layout = compute_layout(tree, plan.manifest.binder)
    assert all(layout.codes[nid] == code for nid, code in plan.toc_codes.items())


# ---- acceptance: every disagreement shown, nothing changed until accepted -------------------

def test_analyze_shows_disagreements_and_changes_nothing(legacy):
    before = _snapshot(legacy)
    plan = importer.analyze(legacy)
    assert _snapshot(legacy) == before
    assert not (legacy / ".binder").exists()

    assert not plan.unmatched_files and not plan.unmatched_nodes
    dis = {(d.entity_id, d.what): d for d in plan.disagreements()}
    assert set(dis) == {("points-of-contact", "header code"), ("meeting-minutes", "filename number")}
    assert dis[("points-of-contact", "header code")].file_says == "A.05"
    assert dis[("points-of-contact", "header code")].toc_says == "A.04"
    assert "TOC wins" in dis[("meeting-minutes", "filename number")].describe()
    # the stale header code did not pair the wrong file
    assert plan.pairs["points-of-contact"].file == f"{ADMIN}/04_Points_of_Contact.docx"
    assert plan.pairs["budget-overview"].file == f"{ADMIN}/05_Budget_Overview.xlsx"
    assert plan.pairs["meeting-minutes"].file == f"{ADMIN}/05_Meeting_Minutes.docx"


def test_accept_writes_only_binder_folder(legacy):
    before = _snapshot(legacy)
    plan = importer.analyze(legacy)
    importer.accept(plan)
    after = _snapshot(legacy)
    outside = {k: v for k, v in after.items() if not k.startswith(".binder/")}
    assert outside == before
    m = Manifest.load(legacy)
    assert m.binder.ref_label == "Packet Ref #"  # the majority label wins over "Ref #:"
    assert m.binder.header_lines == ["{{org_short}}", "North Star Media", "Management Binder", "{{doc_title}}"]
    assert m.binder.filenames.toc == "TOC_m"
    tree = Tree(m)
    assert tree.find("open-points-action-items").get("filename") == "Open_Points"
    assert tree.find("control-section").get("filename") == "Cover_Sheet_CS"
    log = json.loads((legacy / ".binder" / "import.json").read_text())
    assert len(log["disagreements"]) == 2
    assert log["previous_headers"][f"{ADMIN}/03_Top_10_List.docx"][0].endswith("Ref #: A.03")
    assert (legacy / ".binder" / "templates" / "Tabs.docx").is_file()  # legacy tab layout kept as template


def test_manual_pairing_override(legacy):
    plan = importer.analyze(legacy)
    plan.set_pair("points-of-contact", None)
    assert f"{ADMIN}/04_Points_of_Contact.docx" in plan.unmatched_files
    plan.set_pair("points-of-contact", f"{ADMIN}/04_Points_of_Contact.docx")
    assert plan.pairs["points-of-contact"].method == "manual"
    plan.set_pair("org-chart" if "org-chart" in plan.pairs else "organizational-chart", f"{ADMIN}/04_Points_of_Contact.docx")
    assert plan.pairs["points-of-contact"].file is None  # a file can only be paired once


# ---- claims ---------------------------------------------------------------------------------

def test_claim_scan_and_partial_approval(legacy):
    importer.accept(importer.analyze(legacy))
    hits = importer.claim_scan(legacy)
    by = {(h.file.split("/")[-1], h.code) for h in hits}
    assert ("01_North_Star_MRP_Summary.docx", "A.03") in by
    assert ("07_Social_Media_Outreach_Efforts.docx", "A.03.PB") in by
    assert ("01_Open_Points.xlsx", "A.03") in by
    assert all(h.target_id for h in hits)

    mrp = legacy / ADMIN / "01_North_Star_MRP_Summary.docx"
    text_before = refs.full_text_docx(mrp)
    approved = [h for h in hits if h.file.endswith("01_North_Star_MRP_Summary.docx") and h.code == "A.03"]
    rest = [h for h in hits if h not in approved]
    r = importer.apply_claims(legacy, approved, rest)
    assert len(r.claimed) == 1 and not r.skipped
    assert refs.full_text_docx(mrp) == text_before  # wrapped, not changed
    controls = {h.target_id: h.current for h in word.scan_refs(Document(str(mrp)), {})}
    assert controls == {"top-10-list": "A.03"}
    log = json.loads((legacy / ".binder" / "import.json").read_text())
    assert len(log["unclaimed"]) == len(rest)  # unapproved hits are listed for later
    assert any(f.name.endswith("_claims") for f in (legacy / ".binder" / "backups").iterdir())


def test_claim_split_across_runs_and_inside_cell_text(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("See (A.")
    p.add_run("03) and ")
    p.add_run("B.01", ).bold = True
    p.add_run(".")
    src = tmp_path / "a.docx"
    doc.save(str(src))
    hits = refs.scan_docx(src, "a.docx", {"A.03": "x-one", "B.01": "y-two"})
    assert [h.code for h in hits] == ["A.03", "B.01"]
    r = refs.claim_docx(src, tmp_path / "b.docx", hits)
    assert len(r.claimed) == 2
    out = Document(str(tmp_path / "b.docx"))
    # python-docx's paragraph.text skips runs inside content controls, so read every w:t
    full = "".join(t.text for t in out.paragraphs[0]._p.iter(qn("w:t")))
    assert full == "See (A.03) and B.01."
    assert {h.target_id: h.current for h in word.scan_refs(out, {})} == {"x-one": "A.03", "y-two": "B.01"}

    assert refs.claim_cell_formula("A.03", [refs.Hit("f", "S!A1", 0, "A.03", "x-one", "")]) == "=bb_x_one"
    assert refs.claim_cell_formula('see A.03 "now"', [refs.Hit("f", "S!A1", 4, "A.03", "x-one", "")]) == \
        '="see "&bb_x_one&" ""now"""'


def test_code_pattern():
    found = [m.group(0) for m in refs.CODE_RE.finditer("A.03, AB.12.01.PB, 0.01, e.g. 1.5, A.3, Z.99x, (D.01.01)")]
    assert found == ["A.03", "AB.12.01.PB", "0.01", "D.01.01"]


# ---- acceptance: after first sync headers are uniform and claimed refs resolve ----------------

def test_first_sync_after_import(legacy):
    plan = importer.analyze(legacy)
    importer.accept(plan)
    hits = importer.claim_scan(legacy)
    importer.apply_claims(legacy, hits)
    text_before = {h.file: refs.full_text_docx(legacy / h.file) for h in hits if h.file.endswith(".docx")}

    p = sync.plan(legacy)
    assert not p.orphans and not p.warnings, p.render()
    moves = {a.path: a.target for a in p.changes if a.kind == sync.MOVE}
    assert moves == {f"{ADMIN}/05_Meeting_Minutes.docx": f"{ADMIN}/06_Meeting_Minutes.docx"}
    r = sync.apply(p)
    assert r.ok, r.failures
    assert sync.plan(legacy).is_empty

    m = Manifest.load(legacy)
    layout = compute_layout(Tree(m), m.binder)
    for e in layout.items():
        path = legacy / e.rel_path
        assert path.is_file(), e.rel_path
        if e.ext == "docx":
            assert word.read_code(path) == e.code, e.rel_path
            assert "Packet Ref # <bb:code>" in word.header_text(path)  # uniform label, incl. the "Ref #:" one
        elif e.ext == "xlsx":
            assert excel.read_code(path) == e.code, e.rel_path
            assert excel.load(path).worksheets[0].oddHeader.right.text.startswith("Packet Ref # ")
    assert word.read_code(legacy / ADMIN / "04_Points_of_Contact.docx") == "A.04"  # stale code fixed

    # claimed references resolve
    mrp = Document(str(legacy / ADMIN / "01_North_Star_MRP_Summary.docx"))
    assert {h.target_id: h.current for h in word.scan_refs(mrp, {})} == {
        "top-10-list": "A.03", "social-media-outreach-efforts": "A.07"}
    op = legacy / "0_Control_Section" / "01_Open_Points.xlsx"
    assert excel.load(op).worksheets[0]["B2"].value == "=bb_top_10_list"
    assert excel.resolve_name(op, "bb_top_10_list") == "A.03"
    # author text intact through import, claims and first sync
    for rel, h in text_before.items():
        cur = rel.replace("05_Meeting", "06_Meeting")
        assert refs.full_text_docx(legacy / cur) == h, rel

    # and the references follow a later move
    s = BinderSession(legacy)
    s.add_document("Press Kit", "docx", "organizational-chart")
    assert s.apply(s.plan()).ok
    mrp = Document(str(legacy / ADMIN / "01_North_Star_MRP_Summary.docx"))
    assert {h.target_id: h.current for h in word.scan_refs(mrp, {})}["top-10-list"] == "A.04"
    assert excel.resolve_name(op, "bb_top_10_list") == "A.04"
    s.close()


# ---- fill from any source (SPEC §9.3) ------------------------------------------------------------

def test_fill_from_file_copies_and_warns_on_drift(binder, legacy, tmp_path):
    s = BinderSession(binder)
    assert s.apply(s.plan()).ok
    node = s.add_document("Borrowed Policy", "docx", "media-policy")
    source = legacy / "B_Policies_and_Procedures" / "01_Media_Policy.docx"
    original = sync.file_hash(source)
    dest = s.fill_from_file(node.id, source)
    assert dest.name == "02_Borrowed_Policy.docx" and dest.is_file()
    assert sync.file_hash(source) == original  # the source is never modified
    cf = s.tree.find(node.id).get("copied_from")
    assert cf["source_path"] == str(source) and cf["source_hash"] == original

    p = s.plan()
    assert not any(a.kind == sync.CREATE and a.entity_id == node.id for a in p.changes)
    assert s.apply(p).ok
    assert word.read_code(dest) == "B.02"  # header taken over
    assert sync.file_hash(source) == original

    with pytest.raises(FileExistsError):
        s.fill_from_file(node.id, source)

    d = Document(str(source))
    d.add_paragraph("A later edit in the source binder.")
    d.save(str(source))
    warns = [a.detail for a in s.plan().warnings]
    assert any("source changed since copy" in w for w in warns)

    s.recopy(node.id)
    assert any("A later edit" in para.text for para in Document(str(dest)).paragraphs)
    assert not any("source changed" in a.detail for a in s.plan().warnings)
    assert any(b.endswith("_fill") for b in s.backups())
    s.close()


def test_fill_rejects_unsupported_files(binder, tmp_path):
    s = BinderSession(binder)
    bad = tmp_path / "notes.txt"
    bad.write_text("x")
    with pytest.raises(ValueError):
        s.fill_from_file("top-10-list", bad)
    s.close()
