"""Phase 1 acceptance tests (SPEC §14): generated documents and the hybrid checklist."""

import shutil

from openpyxl import load_workbook

from binder_builder.engine import excel, sync
from binder_builder.engine.generate import checklist

from .dump import dump
from .fixtures.make_golden import GENERATED_FILES, GOLDEN, golden_name
from .helpers import do_sync, edit_tree

ADMIN = "A_Administration"


# ---- acceptance: generated files match golden fixtures ----------------------------

def test_generated_files_match_golden(binder):
    do_sync(binder)
    for rel in GENERATED_FILES:
        expected = (GOLDEN / golden_name(rel)).read_text(encoding="utf-8")
        assert dump(binder / rel) == expected, f"{rel} differs from golden; rerun make_golden.py if intended"


def test_generated_docs_follow_reorder(binder):
    do_sync(binder)
    m, tree = edit_tree(binder)
    tree.move("campaign-management", None, 1)  # Campaign Management becomes A
    m.save()
    p = sync.plan(binder)
    regen = {a.path for a in p.changes if a.kind == sync.REGENERATE}
    assert "01_TOC_m.xlsx" in regen and "Tabs.docx" in regen
    do_sync(binder)
    toc = dump(binder / "01_TOC_m.xlsx")
    assert "B8='A' @  C8='Campaign Management'" in toc
    assert "A) Campaign Management" in dump(binder / "A_Campaign_Management/00_Cover.docx")
    tabs = dump(binder / "Tabs.docx")
    assert tabs.index("A) Campaign Management") < tabs.index("B) Administration")
    assert sync.plan(binder).is_empty


# ---- acceptance: CKL user columns survive three syncs with reorders ------------------

def _ckl_rows(path):
    wb = load_workbook(path)
    ws = wb["Master Checklist"]
    heads = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(heads)}
    rows = {}
    extras = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col["bb_id"]]:
            rows[r[col["bb_id"]]] = r
        elif any(v is not None for v in r):
            extras.append(r)
    return col, rows, extras


def test_ckl_user_columns_survive_three_syncs_with_reorders(binder):
    do_sync(binder)
    ckl = binder / "02_CKL.xlsx"

    # The user fills in their columns, adds a free row, and re-sorts the sheet.
    wb = load_workbook(ckl)
    ws = wb["Master Checklist"]
    heads = {c.value: c.column for c in ws[1]}
    filled = {}
    for r in range(2, ws.max_row + 1):
        eid = ws.cell(r, heads["bb_id"]).value
        if eid in ("top-10-list", "web-index", "time-sheet", "#toc", "outreach-efforts"):
            ws.cell(r, heads["Hours"], r * 1.5)
            ws.cell(r, heads["Reviewer"], f"Reviewer for {eid}")
            ws.cell(r, heads["Final Review"], "Done")
            filled[eid] = (r * 1.5, f"Reviewer for {eid}", "Done")
    ws.append([None, "My own note row", 9, "Me"])
    data = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    data.reverse()  # user re-sorts
    for i, values in enumerate(data, start=2):
        for j, v in enumerate(values, start=1):
            ws.cell(i, j).value = v  # ws.cell(i, j, None) would not clear
    wb.save(ckl)

    reorders = [
        lambda t: t.move("policies", None, 1),
        lambda t: t.move("top-10-list", "campaign-management", 0),
        lambda t: t.insert_after("org-chart", {"kind": "item", "title": "Press Kit", "ext": "docx"}),
    ]
    for reorder in reorders:
        m, tree = edit_tree(binder)
        reorder(tree)
        m.save()
        do_sync(binder)
        col, rows, extras = _ckl_rows(ckl)
        for eid, (hours, reviewer, final) in filled.items():
            row = rows[eid]
            assert (row[col["Hours"]], row[col["Reviewer"]], row[col["Final Review"]]) == (hours, reviewer, final), eid
            assert row[col["Reference"]] == f"={excel.name_for(eid)}"
        assert any(e[col["Description"]] == "My own note row" for e in extras)
        assert sync.plan(binder).is_empty

    # the new document got exactly one row, placed after its preceding sibling
    col, rows, _ = _ckl_rows(ckl)
    assert rows["press-kit"][col["Description"]] == "Press Kit"
    wb = load_workbook(ckl)
    ids = [c.value for c in wb["Master Checklist"]["G"]]
    assert ids.index("press-kit") == ids.index("org-chart") + 1


def test_ckl_removed_document_row_kept_and_flagged(binder):
    do_sync(binder)
    m, tree = edit_tree(binder)
    tree.remove("org-chart")
    m.save()
    p = sync.plan(binder)
    assert any(a.kind == sync.CHECKLIST and "org-chart" in a.detail for a in p.changes)
    do_sync(binder)
    wb = load_workbook(binder / "02_CKL.xlsx")
    ws = wb["Master Checklist"]
    row = next(r for r in ws.iter_rows(min_row=2) if r[6].value == "org-chart")
    assert row[0].value == checklist.REMOVED
    assert row[0].fill.fgColor.rgb == checklist.REMOVED_RGB
    assert sync.plan(binder).is_empty


def test_ckl_with_user_formulas_appends_instead_of_inserting(binder):
    do_sync(binder)
    ckl = binder / "02_CKL.xlsx"
    wb = load_workbook(ckl)
    ws = wb["Master Checklist"]
    total_row = ws.max_row + 2
    ws.cell(total_row, 3, f"=SUM(C2:C{ws.max_row})")
    wb.save(ckl)
    m, tree = edit_tree(binder)
    tree.insert_after("org-chart", {"kind": "item", "title": "Press Kit", "ext": "docx"})
    m.save()
    p = sync.plan(binder)
    assert any("at the end" in a.detail for a in p.changes if a.kind == sync.CHECKLIST)
    do_sync(binder)
    ws = load_workbook(ckl)["Master Checklist"]
    assert ws.cell(total_row, 3).value.startswith("=SUM(C2:C")  # user formula untouched
    assert ws.cell(total_row + 1, 7).value == "press-kit"


def test_missing_ckl_is_created_from_template(binder):
    (binder / "02_CKL.xlsx").rename(binder / "old_ckl.xlsx")
    do_sync(binder)
    col, rows, _ = _ckl_rows(binder / "02_CKL.xlsx")
    assert "top-10-list" in rows and "#toc" in rows
    assert (binder / "old_ckl.xlsx").exists()


# ---- acceptance: tabs warn on a 40-character label ---------------------------------

def test_tabs_warn_on_40_character_label(binder):
    label = "Community Partnerships and Media Liaison"
    assert len(label) == 40
    m, tree = edit_tree(binder)
    tree.insert({"kind": "section", "label": label}, None, 3)
    m.save()
    p = sync.plan(binder)
    warns = [a.detail for a in p.warnings]
    assert any(label in w and "over the fit limit" in w for w in warns), warns
    do_sync(binder)
    assert label in dump(binder / "Tabs.docx")


def test_tabs_overflow_appends_table_copies(binder):
    m, tree = edit_tree(binder)
    for i in range(5):
        tree.insert({"kind": "section", "label": f"Extra {i}"})
    m.save()
    do_sync(binder)
    tabs = dump(binder / "Tabs.docx")
    assert tabs.count("<table>") == 2
    assert "<page break>" in tabs
    assert "J) Extra 4" in tabs


# ---- templates ------------------------------------------------------------------------

def test_binder_template_folder_overrides_defaults(binder):
    do_sync(binder)
    from docx import Document
    from binder_builder.engine.generate import DEFAULT_TEMPLATES
    tdir = binder / ".binder" / "templates"
    tdir.mkdir()
    doc = Document(str(DEFAULT_TEMPLATES / "Section_Cover.docx"))
    doc.paragraphs[0].insert_paragraph_before("CUSTOM SECTION COVER")
    doc.save(str(tdir / "Section_Cover.docx"))
    p = sync.plan(binder)
    regen = [a for a in p.changes if a.kind == sync.REGENERATE]
    assert len(regen) == 6 and all("template changed" in a.detail for a in regen)
    assert not p.orphans  # the template folder is the app's, not orphaned files
    do_sync(binder)
    assert "CUSTOM SECTION COVER" in dump(binder / ADMIN / "00_Cover.docx")
    assert sync.plan(binder).is_empty


def test_continuous_mode_adds_page_column_heading(binder):
    m, _ = edit_tree(binder)
    m.raw["binder"]["page_numbers"] = "continuous"
    m.save()
    do_sync(binder)
    assert "E3='Page'" in dump(binder / "01_TOC_m.xlsx")


def test_regenerated_docs_are_backed_up(binder):
    r = do_sync(binder)
    backup = binder / ".binder" / "backups" / r.backup
    assert (backup / "01_TOC_m.xlsx").is_file()
    assert (backup / "Tabs.docx").is_file()
    assert (backup / "00_Cover.docx").is_file()
