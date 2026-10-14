import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from binder_builder.engine import word
from binder_builder.engine.codes import HeaderSpec

from .conftest import FIXTURE

SPEC = HeaderSpec(("North Star", "North Star Media", "Management Binder", "Top 10 List"), "A.03",
                  "Packet Ref #", ("", "", ""))


def test_stamp_header_then_matches(tmp_path):
    src = FIXTURE / "A_Administration/03_Top_10_List.docx"
    dst = tmp_path / "out.docx"
    assert not word.header_matches(Document(str(src)), SPEC)
    word.update(src, dst, SPEC, {})
    doc = Document(str(dst))
    assert word.header_matches(doc, SPEC)
    assert word.read_code(dst) == "A.03"
    text = word.header_text(dst)
    assert "Packet Ref # <bb:code>A.03</>" in text
    assert "{PAGE}" in text and "{NUMPAGES}" in text
    assert word.body_hash(src) == word.body_hash(dst)


def test_first_page_header_written_too(tmp_path):
    src = FIXTURE / "A_Administration/06_Meeting_Minutes.docx"
    dst = tmp_path / "out.docx"
    word.update(src, dst, SPEC, {})
    doc = Document(str(dst))
    first = doc.sections[0].first_page_header
    assert "bb:code" in word._part_signature(first._element)


def test_signature_survives_word_style_complex_fields(tmp_path):
    """Word rewrites fldSimple as complex fields on save; that must not trigger a restamp."""
    src = FIXTURE / "A_Administration/02_Organizational_Chart.docx"
    dst = tmp_path / "out.docx"
    word.update(src, dst, SPEC, {})
    doc = Document(str(dst))
    ftr = doc.sections[0].footer._element
    for fld in list(ftr.iter(qn("w:fldSimple"))):
        instr = fld.get(qn("w:instr"))
        parent = fld.getparent()
        idx = parent.index(fld)
        runs = []
        for kind in ("begin", None, "separate", "result", "end"):
            r = OxmlElement("w:r")
            if kind in ("begin", "separate", "end"):
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                r.append(fc)
            elif kind is None:
                it = OxmlElement("w:instrText")
                it.text = instr
                r.append(it)
            else:
                t = OxmlElement("w:t")
                t.text = "7"
                r.append(t)
            runs.append(r)
        parent.remove(fld)
        for k, r in enumerate(runs):
            parent.insert(idx + k, r)
    assert word.header_matches(doc, SPEC)


def test_body_hash_ignores_bb_controls_only(tmp_path):
    src = FIXTURE / "A_Administration/01_North_Star_MRP_Summary.docx"
    h0 = word.body_hash(src)
    doc = Document(str(src))
    assert word.update_refs(doc, {"top-10-list": "A.09", "outreach-efforts": "A.10"}) == 2
    p = tmp_path / "refs.docx"
    doc.save(str(p))
    assert word.body_hash(p) == h0
    doc.add_paragraph("an author's new sentence")
    doc.save(str(p))
    assert word.body_hash(p) != h0


def test_refs_scan_and_broken():
    doc = Document(str(FIXTURE / "A_Administration/01_North_Star_MRP_Summary.docx"))
    hits = word.scan_refs(doc, {"top-10-list": "A.03"})
    by_id = {h.target_id: h for h in hits}
    assert by_id["top-10-list"].current == "A.03" and by_id["top-10-list"].desired == "A.03"
    assert by_id["outreach-efforts"].desired == word.BROKEN
    assert "outreach summary" in by_id["outreach-efforts"].context


def test_locked_spec_none_leaves_header(tmp_path):
    src = FIXTURE / "B_Policies_and_Procedures/03_Licensed_Release_Form.docx"
    dst = tmp_path / "out.docx"
    word.update(src, dst, None, {})
    assert word.header_text(src) == word.header_text(dst)


def test_create_blank(tmp_path):
    p = tmp_path / "new.docx"
    word.create_blank(p)
    word.update(p, p, SPEC, {})
    assert word.read_code(p) == "A.03"
