"""Phase 4 acceptance tests (SPEC §14): outputs.

"Merged PDF has TOC as page 1, correct bookmarks, correct running numbers; print
manifest lists exactly the files changed since the last print."

A fake exporter stands in for Word/Excel so page counts are known exactly;
test_office_smoke.py runs the real COM exporter.
"""

from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from binder_builder.engine import pdfdoc, sync
from binder_builder.output import merge_pdf, print_manifest
from binder_builder.output.merge_pdf import NeedsSync, OutputOptions

from .dump import dump
from .helpers import do_sync, edit_tree

TWO_PAGE = {"01_North_Star_MRP_Summary.docx", "05_Budget_Overview.xlsx", "01_TOC_m.xlsx"}


class FakeExporter:
    """Makes a PDF with a known page count; each page says which file and page it is."""

    def __init__(self):
        self.calls = []

    def export(self, src: Path, dest: Path) -> None:
        self.calls.append(src.name)
        dest.parent.mkdir(parents=True, exist_ok=True)
        if src.suffix == ".pdf":
            dest.write_bytes(src.read_bytes())
            return
        n = 2 if src.name in TWO_PAGE else 1
        c = canvas.Canvas(str(dest), pagesize=letter, invariant=1)
        for i in range(1, n + 1):
            c.drawString(72, 700, f"{src.name} p{i}")
            c.showPage()
        c.save()


def _page_texts(path):
    return [p.extract_text() or "" for p in PdfReader(str(path)).pages]


def _outline(reader, items=None, depth=0, out=None):
    out = [] if out is None else out
    for it in reader.outline if items is None else items:
        if isinstance(it, list):
            _outline(reader, it, depth + 1, out)
        else:
            out.append((depth, it.title, reader.get_destination_page_number(it)))
    return out


@pytest.fixture
def continuous(binder):
    m, _ = edit_tree(binder)
    m.raw["binder"]["page_numbers"] = "continuous"
    m.save()
    do_sync(binder)
    return binder


# ---- acceptance: merged PDF ----------------------------------------------------------

def test_merged_pdf_toc_is_page_1_with_bookmarks_and_running_numbers(continuous):
    b = continuous
    r = merge_pdf.export_binder(b, OutputOptions(merged=True), FakeExporter())
    reader = PdfReader(str(r.merged))
    texts = _page_texts(r.merged)

    # binder cover first and unnumbered; the TOC is page 1
    assert "00_Cover.docx p1" in texts[0] and texts[0].strip().splitlines()[-1] != "1"
    assert "01_TOC_m.xlsx p1" in texts[1]
    assert reader.page_labels[:4] == ["Cover", "1", "2", "3"]

    # running numbers bottom-center on every page from the TOC onward
    for i, t in enumerate(texts[1:], start=1):
        assert t.strip().endswith(str(i)), (i, t)
    included = [e for e in merge_pdf.printable(sync.plan(b).layout)]
    assert len(texts) == sum(r.pages[e.entity_id] for e in included)

    # bookmarks: every section and document, nested, pointing at the right page
    marks = _outline(reader)
    titles = [t for _, t, _ in marks]
    assert titles[:3] == ["Cover", "Table of Contents", "Master Checklist"]
    assert (0, "A) Administration", r.start_pages["administration#cover"]) in marks
    assert (1, "A.03 Top 10 List", r.start_pages["top-10-list"]) in marks
    assert (0, "Content Mediums", r.start_pages["website-media#cover"]) in marks
    assert (1, "D) Website Media", r.start_pages["website-media#cover"]) in marks
    assert (3, "D.01.01 Weekly Digest Monthly Summary", r.start_pages["weekly-digest"]) in marks
    assert "B.04 Retired Policy" not in titles  # included: false stays out of the binder
    assert len(marks) == len(included) + 1  # plus the "Content Mediums" heading
    # each bookmark's page shows that document's first page
    for depth, title, page in marks:
        if title.startswith("A.03 "):
            assert "03_Top_10_List.docx p1" in texts[page]


def test_toc_page_column_filled_from_real_page_counts(continuous):
    b = continuous
    r = merge_pdf.export_binder(b, OutputOptions(merged=True), FakeExporter())
    toc = dump(b / "01_TOC_m.xlsx")
    # A.01 (MRP summary) starts on its running page; the page after it starts A.02
    a01 = r.start_pages["mrp-summary"]
    assert f"B9='A.01' @  C9='North Star MRP Summary'  E9='{a01}'" in toc or f"E9={a01}" in toc
    assert r.start_pages["org-chart"] == a01 + 2  # MRP summary is two pages
    assert "E3='Page'" in toc
    assert sync.plan(b).is_empty  # the regenerated TOC is what sync expects


def test_per_document_mode_has_no_running_numbers(binder):
    do_sync(binder)
    r = merge_pdf.export_binder(binder, OutputOptions(merged=True), FakeExporter())
    texts = _page_texts(r.merged)
    assert all(not t.strip().splitlines()[-1].strip().isdigit() for t in texts)
    assert "Page" not in dump(binder / "01_TOC_m.xlsx").split("E3=")[-1][:8]


def test_merged_needs_a_synced_binder(binder):
    with pytest.raises(NeedsSync):
        merge_pdf.export_binder(binder, OutputOptions(merged=True), FakeExporter())


def test_per_document_pdfs_and_export_cache(binder):
    do_sync(binder)
    ex = FakeExporter()
    r = merge_pdf.export_binder(binder, OutputOptions(), ex)
    assert (binder / ".binder" / "pdf" / "A_Administration" / "03_Top_10_List.pdf").is_file()
    assert r.pdfs["top-10-list"].name == "03_Top_10_List.pdf"
    first = len(ex.calls)
    ex2 = FakeExporter()
    merge_pdf.export_binder(binder, OutputOptions(), ex2)
    assert ex2.calls == []  # nothing changed: nothing re-exported
    d = Document(str(binder / "A_Administration" / "03_Top_10_List.docx"))
    d.add_paragraph("edit")
    d.save(str(binder / "A_Administration" / "03_Top_10_List.docx"))
    merge_pdf.export_binder(binder, OutputOptions(), ex2)
    assert ex2.calls == ["03_Top_10_List.docx"]
    assert first > 20


# ---- acceptance: print manifest lists exactly the changed files --------------------------

def test_print_manifest_lists_exactly_changed_files(binder):
    do_sync(binder)
    ex = FakeExporter()
    r = merge_pdf.export_binder(binder, OutputOptions(), ex)
    entries = print_manifest.read(binder)
    assert len(entries) == len(r.documents) and all(e["changed_since_last_print"] for e in entries)
    assert set(entries[0]) == {"path", "pdf_path", "code", "title", "pages", "changed_since_last_print"}
    assert entries[1]["code"] == "TOC" and entries[1]["pages"] == 2

    print_manifest.mark_printed(binder, r.documents)
    r = merge_pdf.export_binder(binder, OutputOptions(scope="changed"), ex)
    assert r.documents == [] and print_manifest.read(binder) == [] or r.manifest is None

    # an author edits one document; a TOC change renumbers two more
    d = Document(str(binder / "A_Administration" / "02_Organizational_Chart.docx"))
    d.add_paragraph("A new line from the author.")
    d.save(str(binder / "A_Administration" / "02_Organizational_Chart.docx"))
    m, tree = edit_tree(binder)
    tree.move("outreach-efforts", "administration", 6)  # swap A.06 and A.07
    m.save()
    do_sync(binder)
    r = merge_pdf.export_binder(binder, OutputOptions(scope="changed"), ex)
    changed = {e["code"] for e in print_manifest.read(binder)}
    # exactly what prints differently: the org chart (edited), the two swapped documents (new
    # headers), the MRP summary (its reference now reads A.06), the TOC, the CKL and cover A.00.
    # Workbooks whose hidden _binder sheet changed but whose printed cells didn't are not listed.
    assert changed == {"A.01", "A.02", "A.06", "A.07", "TOC", "CKL", "A.00"}
    assert all(e["changed_since_last_print"] for e in print_manifest.read(binder))


def test_selected_scope(binder):
    do_sync(binder)
    r = merge_pdf.export_binder(binder, OutputOptions(scope="selected", selected=["top-10-list", "web-index"]),
                                FakeExporter())
    assert r.documents == ["top-10-list", "web-index"]
    assert [e["code"] for e in print_manifest.read(binder)] == ["A.03", "D.01"]


# ---- PDF items get a header overlay that never stacks --------------------------------------

def test_pdf_item_header_overlay(binder):
    rel = "B_Policies_and_Procedures/02_Signed_Policy_Acknowledgement.pdf"
    original = (binder / rel).read_bytes()
    do_sync(binder)
    stash = pdfdoc.source_path(binder, "signed-acknowledgement")
    assert stash.read_bytes() == original
    text = _page_texts(binder / rel)[0]
    assert "Packet Ref # B.02" in text and "Page 1 of 1" in text
    assert "Signed Policy Acknowledgement (placeholder scan)" in text
    assert sync.plan(binder).is_empty

    # renumber: overlay regenerated from the original, never stacked
    m, tree = edit_tree(binder)
    tree.move("signed-acknowledgement", "policies", 0)
    m.save()
    do_sync(binder)
    text = _page_texts(binder / "B_Policies_and_Procedures/01_Signed_Policy_Acknowledgement.pdf")[0]
    assert text.count("Packet Ref #") == 1 and "B.01" in text and "B.02" not in text

    # the user drops in a new scan: it becomes the new original
    new = binder / "B_Policies_and_Procedures/01_Signed_Policy_Acknowledgement.pdf"
    c = canvas.Canvas(str(new), pagesize=letter, invariant=1)
    for i in range(2):
        c.drawString(72, 700, f"New scan page {i + 1}")
        c.showPage()
    c.save()
    p = sync.plan(binder)
    assert any("original kept" in a.detail for a in p.changes)
    assert sync.apply(p).ok
    texts = _page_texts(new)
    assert len(texts) == 2 and "Page 2 of 2" in texts[1] and "New scan page 2" in texts[1]
    assert sync.plan(binder).is_empty


def test_locked_pdf_is_left_alone(binder):
    m, tree = edit_tree(binder)
    tree.find("signed-acknowledgement").set("locked_header", True)
    m.save()
    rel = "B_Policies_and_Procedures/02_Signed_Policy_Acknowledgement.pdf"
    before = (binder / rel).read_bytes()
    do_sync(binder)
    assert (binder / rel).read_bytes() == before
